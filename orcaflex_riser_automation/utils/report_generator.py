"""
Report Generator
================
Generates Excel and text reports for riser installation
static analysis results.
"""

import os
import logging
from datetime import datetime
from typing import Optional

logger = logging.getLogger(__name__)


# ─────────────────────────────────────────────────────────────
# Optional imports
# ─────────────────────────────────────────────────────────────
try:
    import openpyxl
    from openpyxl.styles import (Font, PatternFill, Alignment,
                                  Border, Side, numbers)
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False
    logger.warning("openpyxl not installed. Excel reports will be skipped. "
                   "Install with: pip install openpyxl")

try:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import matplotlib.gridspec as gridspec
    MATPLOTLIB_AVAILABLE = True
except ImportError:
    MATPLOTLIB_AVAILABLE = False
    logger.warning("matplotlib not installed. Plot generation will be skipped. "
                   "Install with: pip install matplotlib")


# ─────────────────────────────────────────────────────────────
# Excel Report
# ─────────────────────────────────────────────────────────────

def generate_excel_report(all_stage_results: list, output_path: str,
                          project_name: str = "Riser Installation") -> Optional[str]:
    """
    Generate a comprehensive Excel report for all installation stages.

    Args:
        all_stage_results: list of dicts, one per stage, containing:
            - stage_info: dict (id, name, description, vessel_position, etc.)
            - results: dict of extracted OrcaFlex results
            - code_checks: dict of DNV code check results
            - submerged_weight: dict
        output_path: path to save the .xlsx file
        project_name: project name for the report header

    Returns:
        Path to the saved file, or None if failed.
    """
    if not OPENPYXL_AVAILABLE:
        logger.warning("openpyxl not available. Skipping Excel report.")
        return None

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb = openpyxl.Workbook()

    # ── Styles ────────────────────────────────────────────────
    header_font = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
    title_font = Font(name="Calibri", bold=True, size=14, color="1F3864")
    normal_font = Font(name="Calibri", size=10)
    pass_fill = PatternFill("solid", fgColor="C6EFCE")
    fail_fill = PatternFill("solid", fgColor="FFC7CE")
    header_fill = PatternFill("solid", fgColor="1F3864")
    subheader_fill = PatternFill("solid", fgColor="2E75B6")
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )

    # ── Summary Sheet ─────────────────────────────────────────
    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_summary.column_dimensions["A"].width = 30
    ws_summary.column_dimensions["B"].width = 20
    ws_summary.column_dimensions["C"].width = 20
    ws_summary.column_dimensions["D"].width = 20
    ws_summary.column_dimensions["E"].width = 20
    ws_summary.column_dimensions["F"].width = 15

    # Title
    ws_summary.merge_cells("A1:F1")
    title_cell = ws_summary["A1"]
    title_cell.value = f"RISER INSTALLATION STATIC ANALYSIS REPORT"
    title_cell.font = Font(name="Calibri", bold=True, size=16, color="1F3864")
    title_cell.alignment = center_align

    ws_summary.merge_cells("A2:F2")
    proj_cell = ws_summary["A2"]
    proj_cell.value = f"Project: {project_name}  |  Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    proj_cell.font = Font(name="Calibri", size=10, italic=True, color="595959")
    proj_cell.alignment = center_align

    ws_summary.row_dimensions[1].height = 30
    ws_summary.row_dimensions[2].height = 18

    # Column headers
    headers = ["Stage", "Description", "Max Tension (kN)",
               "Min Bend Radius (m)", "Max Von Mises (MPa)", "Status"]
    row = 4
    for col, h in enumerate(headers, 1):
        cell = ws_summary.cell(row=row, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    # Data rows
    for stage_data in all_stage_results:
        row += 1
        info = stage_data.get("stage_info", {})
        checks = stage_data.get("code_checks", {})
        overall_pass = checks.get("overall", {}).get("pass", True)
        fill = pass_fill if overall_pass else fail_fill

        tension_check = checks.get("tension", {})
        br_check = checks.get("bend_radius", {})
        vm_check = checks.get("von_mises", {})

        row_data = [
            f"Stage {info.get('stage_id', '?')}: {info.get('name', '')}",
            info.get("description", ""),
            tension_check.get("max_tension_kN", "N/A"),
            br_check.get("min_bend_radius_m", "N/A"),
            vm_check.get("max_von_mises_MPa", "N/A"),
            "PASS ✓" if overall_pass else "FAIL ✗",
        ]
        for col, val in enumerate(row_data, 1):
            cell = ws_summary.cell(row=row, column=col, value=val)
            cell.font = normal_font
            cell.fill = fill
            cell.alignment = center_align
            cell.border = thin_border

    # ── Per-Stage Sheets ──────────────────────────────────────
    for stage_data in all_stage_results:
        info = stage_data.get("stage_info", {})
        results = stage_data.get("results", {})
        checks = stage_data.get("code_checks", {})
        sw = stage_data.get("submerged_weight", {})

        sheet_name = f"Stage{info.get('stage_id', '?')}"
        ws = wb.create_sheet(title=sheet_name)
        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 18

        # Stage header
        ws.merge_cells("A1:B1")
        ws["A1"].value = f"Stage {info.get('stage_id')}: {info.get('name')}"
        ws["A1"].font = Font(name="Calibri", bold=True, size=13, color="1F3864")
        ws["A1"].alignment = center_align

        ws["A2"].value = "Description:"
        ws["A2"].font = Font(bold=True)
        ws["B2"].value = info.get("description", "")

        ws["A3"].value = "Vessel Position (m):"
        ws["B3"].value = info.get("vessel_position", "")
        ws["A4"].value = "Riser Length Deployed (m):"
        ws["B4"].value = info.get("riser_length_deployed", "")
        ws["A5"].value = "Top Tension (kN):"
        ws["B5"].value = info.get("top_tension", "")

        # Code checks section
        row = 7
        ws.merge_cells(f"A{row}:B{row}")
        ws[f"A{row}"].value = "CODE CHECK RESULTS (DNV-ST-F101)"
        ws[f"A{row}"].font = Font(bold=True, color="FFFFFF")
        ws[f"A{row}"].fill = subheader_fill
        ws[f"A{row}"].alignment = center_align

        row += 1
        for check_name, check_data in checks.items():
            if check_name == "overall":
                continue
            ws[f"A{row}"].value = check_name.replace("_", " ").title()
            ws[f"A{row}"].font = Font(bold=True)
            for key, val in check_data.items():
                row += 1
                ws[f"A{row}"].value = f"  {key}"
                ws[f"B{row}"].value = str(val)
                if key == "pass":
                    ws[f"B{row}"].fill = pass_fill if val else fail_fill
            row += 1

        # Overall status
        overall = checks.get("overall", {})
        ws[f"A{row}"].value = "OVERALL STATUS"
        ws[f"A{row}"].font = Font(bold=True, size=12)
        ws[f"B{row}"].value = overall.get("status", "N/A")
        ws[f"B{row}"].font = Font(bold=True, size=12)
        ws[f"B{row}"].fill = pass_fill if overall.get("pass", True) else fail_fill

        # Submerged weight section
        row += 2
        ws[f"A{row}"].value = "SUBMERGED WEIGHT BREAKDOWN"
        ws[f"A{row}"].font = Font(bold=True, color="FFFFFF")
        ws[f"A{row}"].fill = subheader_fill
        row += 1
        for key, val in sw.items():
            ws[f"A{row}"].value = key.replace("_", " ").title()
            ws[f"B{row}"].value = val
            row += 1

        # Results table
        if results:
            row += 1
            ws[f"A{row}"].value = "RESULTS ALONG ARC LENGTH"
            ws[f"A{row}"].font = Font(bold=True, color="FFFFFF")
            ws[f"A{row}"].fill = header_fill
            row += 1

            variables = list(results.keys())
            for col, var in enumerate(variables, 1):
                cell = ws.cell(row=row, column=col, value=var)
                cell.font = Font(bold=True)
                cell.fill = subheader_fill
                cell.font = Font(bold=True, color="FFFFFF")
                ws.column_dimensions[get_column_letter(col)].width = 20

            n_points = len(results.get("ArcLength", []))
            for i in range(n_points):
                row += 1
                for col, var in enumerate(variables, 1):
                    vals = results.get(var, [])
                    val = vals[i] if i < len(vals) else ""
                    ws.cell(row=row, column=col, value=val)

    wb.save(output_path)
    logger.info(f"Excel report saved: {output_path}")
    return output_path


# ─────────────────────────────────────────────────────────────
# Plot Generation
# ─────────────────────────────────────────────────────────────

def generate_plots(stage_data: dict, output_dir: str) -> list:
    """
    Generate result plots for a single installation stage.

    Args:
        stage_data: dict with stage_info, results, code_checks
        output_dir: directory to save plot images

    Returns:
        List of saved plot file paths.
    """
    if not MATPLOTLIB_AVAILABLE:
        logger.warning("matplotlib not available. Skipping plots.")
        return []

    os.makedirs(output_dir, exist_ok=True)
    info = stage_data.get("stage_info", {})
    results = stage_data.get("results", {})
    stage_id = info.get("stage_id", "?")
    stage_name = info.get("name", "Stage")

    arc = results.get("ArcLength", [])
    if not arc:
        return []

    saved_plots = []

    # ── Plot 1: Structural Results ─────────────────────────────
    fig, axes = plt.subplots(2, 2, figsize=(14, 10))
    fig.suptitle(f"Stage {stage_id}: {stage_name} — Structural Results",
                 fontsize=14, fontweight="bold", color="#1F3864")

    plot_vars = [
        ("Effective Tension", "Effective Tension (kN)", "steelblue", axes[0, 0]),
        ("Bend Moment", "Bend Moment (kN·m)", "darkorange", axes[0, 1]),
        ("Curvature", "Curvature (1/m)", "green", axes[1, 0]),
        ("Von Mises Stress", "Von Mises Stress (MPa)", "crimson", axes[1, 1]),
    ]

    for var, ylabel, color, ax in plot_vars:
        vals = results.get(var, [])
        if vals:
            ax.plot(arc[:len(vals)], vals, color=color, linewidth=1.8)
            ax.set_xlabel("Arc Length (m)", fontsize=9)
            ax.set_ylabel(ylabel, fontsize=9)
            ax.set_title(var, fontsize=10, fontweight="bold")
            ax.grid(True, alpha=0.3)
            ax.tick_params(labelsize=8)
        else:
            ax.text(0.5, 0.5, "No Data", ha="center", va="center",
                    transform=ax.transAxes, fontsize=12, color="gray")
            ax.set_title(var, fontsize=10)

    plt.tight_layout(rect=[0, 0, 1, 0.95])
    plot1_path = os.path.join(output_dir, f"stage{stage_id}_structural.png")
    plt.savefig(plot1_path, dpi=150, bbox_inches="tight")
    plt.close()
    saved_plots.append(plot1_path)
    logger.info(f"Structural plot saved: {plot1_path}")

    # ── Plot 2: Riser Profile (x-z) ───────────────────────────
    x_vals = results.get("x", [])
    z_vals = results.get("z", [])

    if x_vals and z_vals:
        fig2, ax2 = plt.subplots(figsize=(12, 5))
        ax2.plot(x_vals[:len(z_vals)], z_vals, color="navy",
                 linewidth=2.0, label="Riser Profile")
        ax2.axhline(y=0, color="skyblue", linestyle="--",
                    linewidth=1.0, label="Sea Surface")
        ax2.fill_between(x_vals[:len(z_vals)], z_vals,
                         min(z_vals) - 10, alpha=0.1, color="navy")
        ax2.set_xlabel("Horizontal Position (m)", fontsize=10)
        ax2.set_ylabel("Depth (m)", fontsize=10)
        ax2.set_title(f"Stage {stage_id}: {stage_name} — Riser Profile",
                      fontsize=12, fontweight="bold", color="#1F3864")
        ax2.legend(fontsize=9)
        ax2.grid(True, alpha=0.3)
        ax2.invert_yaxis()

        plot2_path = os.path.join(output_dir, f"stage{stage_id}_profile.png")
        plt.savefig(plot2_path, dpi=150, bbox_inches="tight")
        plt.close()
        saved_plots.append(plot2_path)
        logger.info(f"Profile plot saved: {plot2_path}")

    return saved_plots


def generate_summary_plot(all_stage_results: list, output_dir: str) -> Optional[str]:
    """
    Generate a summary comparison plot across all installation stages.
    """
    if not MATPLOTLIB_AVAILABLE:
        return None

    os.makedirs(output_dir, exist_ok=True)

    stage_names = []
    max_tensions = []
    min_bend_radii = []
    max_vm_stresses = []
    statuses = []

    for sd in all_stage_results:
        info = sd.get("stage_info", {})
        checks = sd.get("code_checks", {})
        stage_names.append(f"S{info.get('stage_id', '?')}\n{info.get('name', '')[:10]}")
        max_tensions.append(checks.get("tension", {}).get("max_tension_kN", 0) or 0)
        min_bend_radii.append(checks.get("bend_radius", {}).get("min_bend_radius_m", 0) or 0)
        max_vm_stresses.append(checks.get("von_mises", {}).get("max_von_mises_MPa", 0) or 0)
        statuses.append(checks.get("overall", {}).get("pass", True))

    colors = ["#2ecc71" if p else "#e74c3c" for p in statuses]

    fig, axes = plt.subplots(1, 3, figsize=(16, 6))
    fig.suptitle("Riser Installation — Stage Comparison Summary",
                 fontsize=14, fontweight="bold", color="#1F3864")

    # Tension
    bars = axes[0].bar(stage_names, max_tensions, color=colors, edgecolor="white", linewidth=0.5)
    axes[0].set_title("Max Effective Tension (kN)", fontweight="bold")
    axes[0].set_ylabel("kN")
    axes[0].grid(axis="y", alpha=0.3)
    axes[0].tick_params(axis="x", labelsize=7)

    # Bend Radius
    bars2 = axes[1].bar(stage_names, min_bend_radii, color=colors, edgecolor="white", linewidth=0.5)
    axes[1].set_title("Min Bend Radius (m)", fontweight="bold")
    axes[1].set_ylabel("m")
    axes[1].grid(axis="y", alpha=0.3)
    axes[1].tick_params(axis="x", labelsize=7)

    # Von Mises
    bars3 = axes[2].bar(stage_names, max_vm_stresses, color=colors, edgecolor="white", linewidth=0.5)
    axes[2].set_title("Max Von Mises Stress (MPa)", fontweight="bold")
    axes[2].set_ylabel("MPa")
    axes[2].grid(axis="y", alpha=0.3)
    axes[2].tick_params(axis="x", labelsize=7)

    # Legend
    from matplotlib.patches import Patch
    legend_elements = [Patch(facecolor="#2ecc71", label="PASS"),
                       Patch(facecolor="#e74c3c", label="FAIL")]
    fig.legend(handles=legend_elements, loc="lower center",
               ncol=2, fontsize=10, frameon=True)

    plt.tight_layout(rect=[0, 0.05, 1, 0.95])
    summary_plot_path = os.path.join(output_dir, "summary_comparison.png")
    plt.savefig(summary_plot_path, dpi=150, bbox_inches="tight")
    plt.close()
    logger.info(f"Summary plot saved: {summary_plot_path}")
    return summary_plot_path


# ─────────────────────────────────────────────────────────────
# Text Report
# ─────────────────────────────────────────────────────────────

def generate_text_report(all_stage_results: list, output_path: str,
                         project_name: str = "Riser Installation") -> str:
    """Generate a plain-text summary report."""
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    lines = []
    sep = "=" * 70
    thin_sep = "-" * 70

    lines.append(sep)
    lines.append(f"  RISER INSTALLATION STATIC ANALYSIS REPORT")
    lines.append(f"  Project: {project_name}")
    lines.append(f"  Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    lines.append(sep)
    lines.append("")

    for sd in all_stage_results:
        info = sd.get("stage_info", {})
        checks = sd.get("code_checks", {})
        sw = sd.get("submerged_weight", {})

        lines.append(thin_sep)
        lines.append(f"  STAGE {info.get('stage_id')}: {info.get('name', '').upper()}")
        lines.append(thin_sep)
        lines.append(f"  Description      : {info.get('description', '')}")
        lines.append(f"  Vessel Position  : {info.get('vessel_position', '')} m")
        lines.append(f"  Length Deployed  : {info.get('riser_length_deployed', '')} m")
        lines.append(f"  Top Tension      : {info.get('top_tension', '')} kN")
        lines.append("")

        lines.append("  CODE CHECKS (DNV-ST-F101):")
        for check_name, check_data in checks.items():
            if check_name == "overall":
                continue
            status = "PASS ✓" if check_data.get("pass", True) else "FAIL ✗"
            lines.append(f"    [{status}] {check_name.replace('_', ' ').title()}")
            for k, v in check_data.items():
                if k != "pass":
                    lines.append(f"           {k}: {v}")

        overall = checks.get("overall", {})
        lines.append("")
        lines.append(f"  >>> OVERALL: {overall.get('status', 'N/A')} <<<")
        lines.append("")

        if sw:
            lines.append("  SUBMERGED WEIGHT:")
            for k, v in sw.items():
                lines.append(f"    {k}: {v}")
        lines.append("")

    lines.append(sep)
    lines.append("  END OF REPORT")
    lines.append(sep)

    with open(output_path, "w") as f:
        f.write("\n".join(lines))

    logger.info(f"Text report saved: {output_path}")
    return output_path
